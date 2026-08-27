"""
VoxKage Microsoft Word Custom MCP Server Module
Handles creation, incremental section-based editing (v1 -> v2 versioning),
native LaTeX math formula compilation (MathML -> OMML via MML2OMML.XSL),
rich typography, and native .docx binary export.
"""

import os
import re
import json
import io
from datetime import datetime
from typing import List, Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

import lxml.etree as ET
import latex2mathml.converter

# --- XSLT MATHML TO OMML TRANSFORM INITIALIZATION ---
XSL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "MML2OMML.XSL")
_xslt_transform = None

def get_xslt_transform():
    global _xslt_transform
    if _xslt_transform is None:
        if os.path.exists(XSL_PATH):
            try:
                xslt_doc = ET.parse(XSL_PATH)
                _xslt_transform = ET.XSLT(xslt_doc)
            except Exception as e:
                print(f"[-] Failed to load local MML2OMML.XSL: {e}")
        
        if _xslt_transform is None:
            # Fallback to standard Office paths on Windows
            office_paths = [
                r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL',
                r'C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL',
                r'C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL',
                r'C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL'
            ]
            for p in office_paths:
                if os.path.exists(p):
                    try:
                        xslt_doc = ET.parse(p)
                        _xslt_transform = ET.XSLT(xslt_doc)
                        break
                    except Exception:
                        pass
    return _xslt_transform


def latex_to_omml_element(latex_str: str, is_display: bool = True) -> Optional[Any]:
    """
    Converts LaTeX formulas into true native Microsoft Word OMML XML elements (<m:oMath>).
    """
    clean_eq = latex_str.strip()
    if is_display:
        clean_eq = re.sub(r'^\$\$|\$\$$', '', clean_eq).strip()
        clean_eq = re.sub(r'^\\\[|\\\]$', '', clean_eq).strip()
    else:
        clean_eq = re.sub(r'^\$|\$$', '', clean_eq).strip()
        clean_eq = re.sub(r'^\\\(|\\\)$', '', clean_eq).strip()

    if not clean_eq:
        return None

    try:
        # 1. Convert LaTeX to standard MathML
        mathml_str = latex2mathml.converter.convert(clean_eq)

        # 2. Transform MathML to OMML using XSLT
        transform = get_xslt_transform()
        if transform is not None:
            dom = ET.fromstring(mathml_str.encode('utf-8'))
            omml_tree = transform(dom)
            omml_bytes = ET.tostring(omml_tree, encoding='utf-8')
            return parse_xml(omml_bytes)
    except Exception as e:
        print(f"[-] LaTeX to OMML transform failed for '{clean_eq[:35]}': {e}")

    return None


class WordDocumentEngine:
    def __init__(self):
        self.documents: Dict[str, List[Dict[str, Any]]] = {}

    def create_document(
        self,
        doc_id: str,
        title: str,
        sections: List[Dict[str, Any]],
        author: Optional[str] = None,
        header_text: Optional[str] = None,
        footer_text: Optional[str] = "Page {page} of {pages}",
        style: Optional[Dict[str, Any]] = None
    ) -> Dict[str, Any]:
        """
        Creates a new version 1 (v1) Word Document structure with dynamic theme styling and optional headers.
        """
        style = style or {}
        doc_state = {
            "version": 1,
            "doc_id": doc_id,
            "title": title,
            "author": author,
            "header_text": header_text,
            "footer_text": footer_text,
            "style": {
                "primary_color": style.get("primary_color", "#1F497D"),
                "secondary_color": style.get("secondary_color", "#4A5568"),
                "font_family": style.get("font_family", "Calibri"),
                "font_size_body": style.get("font_size_body", 11),
                "table_header_bg": style.get("table_header_bg", style.get("primary_color", "#1F497D")),
                "table_header_color": style.get("table_header_color", "#FFFFFF"),
                "table_alternate_bg": style.get("table_alternate_bg", "#F5F0E8")
            },
            "created_at": datetime.now().strftime("%Y-%m-%d"),
            "sections": sections
        }
        
        if doc_id not in self.documents:
            self.documents[doc_id] = []
        
        self.documents[doc_id].append(doc_state)
        return doc_state

    def edit_document(
        self,
        doc_id: str,
        target_heading_or_index: str,
        action: str,  # 'replace', 'append', 'insert_after', 'delete'
        new_content: Dict[str, Any]
    ) -> Dict[str, Any]:
        """
        Performs an incremental edit without overwriting unrelated sections,
        producing a new version (e.g. v2) while retaining v1 in history.
        """
        history = self.documents.get(doc_id, [])
        if not history:
            latest = {
                "version": 1,
                "doc_id": doc_id,
                "title": "Untitled Word Document",
                "author": None,
                "header_text": None,
                "footer_text": "Page {page} of {pages}",
                "style": {
                    "primary_color": "#1F497D",
                    "secondary_color": "#4A5568",
                    "font_family": "Calibri",
                    "font_size_body": 11,
                    "table_header_bg": "#1F497D",
                    "table_header_color": "#FFFFFF",
                    "table_alternate_bg": "#F5F0E8"
                },
                "sections": []
            }
        else:
            latest = json.loads(json.dumps(history[-1]))

        next_version = latest["version"] + 1
        latest["version"] = next_version

        sections = latest.get("sections", [])
        modified = False

        for idx, sec in enumerate(sections):
            heading = sec.get("heading", "").strip().lower()
            if target_heading_or_index.strip().lower() in heading or str(idx) == target_heading_or_index.strip():
                if action == "replace":
                    sections[idx] = new_content
                elif action in ("append", "insert_after"):
                    sections.insert(idx + 1, new_content)
                elif action == "delete":
                    sections.pop(idx)
                modified = True
                break

        if not modified:
            sections.append(new_content)

        latest["sections"] = sections
        history.append(latest)
        self.documents[doc_id] = history
        return latest

    def _hex_to_rgb(self, hex_str: str, default: str = "1F497D") -> RGBColor:
        clean = str(hex_str).lstrip("#").strip()
        if len(clean) == 3:
            clean = "".join([c * 2 for c in clean])
        if len(clean) != 6:
            clean = default.lstrip("#")
        try:
            r = int(clean[0:2], 16)
            g = int(clean[2:4], 16)
            b = int(clean[4:6], 16)
            return RGBColor(r, g, b)
        except Exception:
            return RGBColor(0x1F, 0x49, 0x7D)

    def _clean_hex(self, hex_str: str, default: str = "1F497D") -> str:
        clean = str(hex_str).lstrip("#").strip()
        if len(clean) == 3:
            clean = "".join([c * 2 for c in clean])
        if len(clean) != 6:
            return default.lstrip("#").upper()
        return clean.upper()

    def _add_math_equation(self, doc, eq_text: str):
        """
        Inserts a centered mathematical equation paragraph using native OpenXML Math (OMML) in Cambria Math.
        """
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)

        omml_elem = latex_to_omml_element(eq_text, is_display=True)
        if omml_elem is not None:
            p._p.append(omml_elem)
        else:
            clean_eq = re.sub(r'^\$\$|\$\$$|^\\\[|\\\]$', '', eq_text.strip())
            r = p.add_run(clean_eq)
            r.font.name = "Cambria Math"
            r.font.size = Pt(11)
            r.font.italic = True

    def _add_formatted_text_runs(self, paragraph, text: str, font_name: str, base_size: int, secondary_rgb: Optional[RGBColor] = None):
        """
        Tokenizes inline markdown: **bold**, *italic*, `code`, and $math$ into styled OpenXML runs / OMML elements.
        """
        pattern = r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$|\\\(.*?\\\))'
        tokens = re.split(pattern, text)

        for token in tokens:
            if not token:
                continue

            if token.startswith("**") and token.endswith("**") and len(token) >= 4:
                r = paragraph.add_run(token[2:-2])
                r.font.name = font_name
                r.font.size = Pt(base_size)
                r.font.bold = True
                if secondary_rgb:
                    r.font.color.rgb = secondary_rgb

            elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
                r = paragraph.add_run(token[1:-1])
                r.font.name = font_name
                r.font.size = Pt(base_size)
                r.font.italic = True
                if secondary_rgb:
                    r.font.color.rgb = secondary_rgb

            elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
                r = paragraph.add_run(token[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(max(8, base_size - 1))

            elif (token.startswith("$") and token.endswith("$") and len(token) >= 2) or (token.startswith(r"\(") and token.endswith(r"\)")):
                omml_inline = latex_to_omml_element(token, is_display=False)
                if omml_inline is not None:
                    paragraph._p.append(omml_inline)
                else:
                    math_txt = token[1:-1] if token.startswith("$") else token[2:-2]
                    r = paragraph.add_run(math_txt)
                    r.font.name = "Cambria Math"
                    r.font.size = Pt(base_size)
                    r.font.italic = True

            else:
                r = paragraph.add_run(token)
                r.font.name = font_name
                r.font.size = Pt(base_size)
                if secondary_rgb:
                    r.font.color.rgb = secondary_rgb

    def export_to_docx_bytes(self, doc_state: Any) -> bytes:
        """
        Generates binary .docx file bytes using python-docx with native OpenXML headers/footers,
        dynamic color themes, OMML math equations, and custom styled tables.
        """
        if isinstance(doc_state, str):
            try:
                doc_state = json.loads(doc_state)
            except Exception:
                doc_state = {"title": "Document", "sections": [{"heading": "Content", "body": doc_state}]}

        if not isinstance(doc_state, dict):
            doc_state = {"title": "Document", "sections": []}

        style_opts = doc_state.get("style") or {}
        primary_hex = style_opts.get("primary_color", "#1F497D")
        secondary_hex = style_opts.get("secondary_color", "#4A5568")
        font_name = style_opts.get("font_family", "Calibri")
        base_size = style_opts.get("font_size_body", 11)
        table_hdr_bg = style_opts.get("table_header_bg", primary_hex)
        table_hdr_fg = style_opts.get("table_header_color", "#FFFFFF")
        table_alt_bg = style_opts.get("table_alternate_bg", "#F5F0E8")

        primary_rgb = self._hex_to_rgb(primary_hex, "1F497D")
        secondary_rgb = self._hex_to_rgb(secondary_hex, "4A5568")
        table_hdr_fg_rgb = self._hex_to_rgb(table_hdr_fg, "FFFFFF")

        doc = docx.Document()

        # Page Margins (1 inch)
        for sec in doc.sections:
            sec.top_margin = Inches(1.0)
            sec.bottom_margin = Inches(1.0)
            sec.left_margin = Inches(1.0)
            sec.right_margin = Inches(1.0)

            # Header — ONLY if header_text is explicitly provided
            header_text = doc_state.get("header_text")
            if header_text and str(header_text).strip():
                sec.different_first_page_header_footer = True
                header_para = sec.header.paragraphs[0]
                header_para.text = ""
                header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                h_run = header_para.add_run(str(header_text).strip())
                h_run.font.name = font_name
                h_run.font.size = Pt(9)
                h_run.font.color.rgb = secondary_rgb
            else:
                sec.header.paragraphs[0].text = ""

            # Footer — parse {page} and {pages} into native OpenXML fields
            footer_text = doc_state.get("footer_text")
            if footer_text and str(footer_text).strip():
                footer_para = sec.footer.paragraphs[0]
                footer_para.text = ""
                footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                footer_str = str(footer_text).strip()

                tokens = re.split(r'(\{page\}|\{pages\})', footer_str)
                for tok in tokens:
                    if tok == "{page}":
                        fld_xml = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
                        footer_para._p.append(fld_xml)
                    elif tok == "{pages}":
                        fld_xml = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
                        footer_para._p.append(fld_xml)
                    elif tok:
                        f_run = footer_para.add_run(tok)
                        f_run.font.name = font_name
                        f_run.font.size = Pt(9)
                        f_run.font.color.rgb = secondary_rgb

        # Document Title (Clean, no fabricated subtitle underneath)
        title_para = doc.add_paragraph()
        title_para.paragraph_format.space_before = Pt(0)
        title_para.paragraph_format.space_after = Pt(14)
        title_run = title_para.add_run(str(doc_state.get("title", "Document")))
        title_run.font.name = font_name
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = primary_rgb
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Render Sections
        sections = doc_state.get("sections") or []
        if not isinstance(sections, list):
            sections = []

        for sec in sections:
            if not isinstance(sec, dict):
                continue

            # Handle explicit section page breaks
            if sec.get("page_break"):
                doc.add_page_break()

            heading = sec.get("heading")
            if heading:
                raw_level = sec.get("level", 1)
                try:
                    level_int = int(raw_level)
                except (ValueError, TypeError):
                    level_int = 1
                level_int = max(1, min(level_int, 3))

                h_para = doc.add_heading(level=level_int)
                h_para.paragraph_format.keep_with_next = True
                h_para.paragraph_format.space_before = Pt(14 if level_int == 1 else 10)
                h_para.paragraph_format.space_after = Pt(4)

                hrun = h_para.add_run(str(heading))
                hrun.font.name = font_name
                hrun.font.bold = True
                if level_int == 1:
                    hrun.font.size = Pt(18)
                    hrun.font.color.rgb = primary_rgb
                elif level_int == 2:
                    hrun.font.size = Pt(14)
                    hrun.font.color.rgb = secondary_rgb
                else:
                    hrun.font.size = Pt(12)
                    hrun.font.color.rgb = secondary_rgb

            # Paragraphs / Body text / Math formulas
            body = sec.get("body", "")
            if body:
                blocks = str(body).split("\n")
                in_math_block = False
                math_buffer = []

                for line in blocks:
                    stripped = line.strip()
                    if not stripped:
                        continue

                    # Single-line display math
                    if stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) >= 4:
                        self._add_math_equation(doc, stripped)
                        continue
                    elif stripped.startswith(r"\[") and stripped.endswith(r"\]"):
                        self._add_math_equation(doc, stripped)
                        continue
                    elif stripped.startswith("$$"):
                        in_math_block = True
                        math_buffer = [stripped.lstrip("$")]
                        continue
                    elif stripped.endswith("$$") and in_math_block:
                        math_buffer.append(stripped.rstrip("$"))
                        self._add_math_equation(doc, " ".join(math_buffer))
                        in_math_block = False
                        math_buffer = []
                        continue
                    elif in_math_block:
                        math_buffer.append(stripped)
                        continue

                    # Bullet lists
                    if stripped.startswith("•") or stripped.startswith("- ") or stripped.startswith("* "):
                        clean_line = re.sub(r'^[•\-\*]\s*', '', stripped)
                        p = doc.add_paragraph(style='List Bullet')
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.15
                        self._add_formatted_text_runs(p, clean_line, font_name, base_size)
                    # Numbered lists
                    elif re.match(r'^\d+[\.\)]\s+', stripped):
                        clean_line = re.sub(r'^\d+[\.\)]\s+', '', stripped)
                        p = doc.add_paragraph(style='List Number')
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.15
                        self._add_formatted_text_runs(p, clean_line, font_name, base_size)
                    # Quotes / Callouts
                    elif stripped.startswith(">"):
                        clean_line = stripped.lstrip("> ").strip()
                        p = doc.add_paragraph()
                        p.paragraph_format.left_indent = Inches(0.4)
                        p.paragraph_format.space_before = Pt(4)
                        p.paragraph_format.space_after = Pt(6)
                        p.paragraph_format.line_spacing = 1.15
                        self._add_formatted_text_runs(p, clean_line, font_name, base_size, secondary_rgb=secondary_rgb)
                    # Regular paragraph with inline math support
                    else:
                        p = doc.add_paragraph()
                        p.paragraph_format.space_after = Pt(6)
                        p.paragraph_format.line_spacing = 1.15
                        self._add_formatted_text_runs(p, stripped, font_name, base_size)

            # Table (if present)
            table_data = sec.get("table")
            if table_data and isinstance(table_data, dict):
                headers = table_data.get("headers") if isinstance(table_data.get("headers"), list) else []
                rows = table_data.get("rows") if isinstance(table_data.get("rows"), list) else []

                if headers or rows:
                    first_row_len = len(rows[0]) if rows and isinstance(rows[0], list) else 1
                    t_cols = len(headers) if headers else first_row_len
                    t_rows = (1 if headers else 0) + len(rows)
                    if t_cols > 0 and t_rows > 0:
                        table = doc.add_table(rows=t_rows, cols=t_cols)
                        table.alignment = WD_TABLE_ALIGNMENT.CENTER

                        tblPr = table._tbl.tblPr
                        try:
                            tblCellMar = parse_xml(
                                r'<w:tblCellMar %s>'
                                r'  <w:top w:w="120" w:type="dxa"/>'
                                r'  <w:left w:w="160" w:type="dxa"/>'
                                r'  <w:bottom w:w="120" w:type="dxa"/>'
                                r'  <w:right w:w="160" w:type="dxa"/>'
                                r'</w:tblCellMar>' % nsdecls('w')
                            )
                            tblPr.append(tblCellMar)
                        except Exception:
                            pass

                        try:
                            tblBorders = parse_xml(
                                r'<w:tblBorders %s>'
                                r'  <w:top w:val="single" w:sz="6" w:space="0" w:color="%s"/>'
                                r'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="%s"/>'
                                r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E6DFD8"/>'
                                r'  <w:left w:val="none"/>'
                                r'  <w:right w:val="none"/>'
                                r'  <w:insideV w:val="none"/>'
                                r'</w:tblBorders>' % (nsdecls('w'), self._clean_hex(table_hdr_bg), self._clean_hex(table_hdr_bg))
                            )
                            tblPr.append(tblBorders)
                        except Exception:
                            pass

                        row_idx = 0
                        if headers:
                            hdr_cells = table.rows[0].cells
                            for c_idx, h_text in enumerate(headers):
                                if c_idx < len(hdr_cells):
                                    cell = hdr_cells[c_idx]
                                    cell.text = str(h_text)
                                    try:
                                        shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), self._clean_hex(table_hdr_bg)))
                                        cell._tc.get_or_add_tcPr().append(shd)
                                        for p in cell.paragraphs:
                                            for r in p.runs:
                                                r.font.name = font_name
                                                r.font.bold = True
                                                r.font.color.rgb = table_hdr_fg_rgb
                                    except Exception:
                                        pass
                            row_idx = 1

                        for r_data_idx, r_data in enumerate(rows):
                            if isinstance(r_data, list) and row_idx < len(table.rows):
                                row_cells = table.rows[row_idx].cells
                                is_alt = (r_data_idx % 2 == 1)
                                for c_idx, cell_value in enumerate(r_data):
                                    if c_idx < len(row_cells):
                                        cell = row_cells[c_idx]
                                        cell.text = str(cell_value)
                                        if is_alt and table_alt_bg:
                                            try:
                                                shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), self._clean_hex(table_alt_bg)))
                                                cell._tc.get_or_add_tcPr().append(shd)
                                            except Exception:
                                                pass
                                        for p in cell.paragraphs:
                                            for r in p.runs:
                                                r.font.name = font_name
                                                r.font.size = Pt(10)
                                row_idx += 1

            doc.add_paragraph()

        target_stream = io.BytesIO()
        doc.save(target_stream)
        target_stream.seek(0)
        return target_stream.getvalue()


# Global Singleton Word Engine instance
word_engine = WordDocumentEngine()


# --- WORD MCP TOOL SCHEMAS FOR AGENT REGISTRATION ---
WORD_TOOLS = [
    {
        "name": "create_word_document",
        "description": "Create a beautifully styled Microsoft Word document (.docx) artifact. Supports LaTeX math formulas ($$ ... $$ and $ ... $), markdown formatting (bold, italics, bullet/numbered lists, quotes), custom typography, color themes, and optional tables.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "doc_id": {"type": "string", "description": "Unique document ID or filename slug without extension (e.g. 'math_compendium')"},
                "title": {"type": "string", "description": "Title of the Word document (e.g. 'Complex Mathematical Formula Compendium')"},
                "sections": {
                    "type": "array",
                    "description": "List of document sections with heading, level (1, 2, or 3), body text with optional LaTeX math, optional table, or page_break",
                    "items": {
                        "type": "object",
                        "properties": {
                            "heading": {"type": "string", "description": "Section heading text"},
                            "level": {"type": "integer", "description": "Heading level 1, 2, or 3"},
                            "body": {"type": "string", "description": "Body text with markdown formatting and LaTeX math formulas ($$ equation $$ for display, $ equation $ for inline)"},
                            "page_break": {"type": "boolean", "description": "Set to true to insert an explicit page break before this section"},
                            "table": {
                                "type": "object",
                                "description": "Optional table. Only include when tabular data is relevant.",
                                "properties": {
                                    "headers": {"type": "array", "items": {"type": "string"}},
                                    "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}}
                                }
                            }
                        }
                    }
                },
                "header_text": {"type": "string", "description": "Optional top running header for pages 2+. Leave empty or omit if no running header is needed."},
                "footer_text": {"type": "string", "description": "Optional footer text format (e.g. 'Page {page} of {pages}'). Default is 'Page {page} of {pages}'."},
                "style": {
                    "type": "object",
                    "description": "Optional document theme & styling parameters tailored to document type",
                    "properties": {
                        "primary_color": {"type": "string", "description": "Primary accent / heading color (hex string e.g. '#1F497D', '#2D3748', '#8C1D40')"},
                        "secondary_color": {"type": "string", "description": "Secondary accent color (hex string e.g. '#4A5568', '#718096')"},
                        "font_family": {"type": "string", "description": "Font family e.g. 'Georgia', 'Calibri', 'Arial', 'Times New Roman', 'Garamond', 'Segoe UI'"},
                        "font_size_body": {"type": "integer", "description": "Body font size in Pt (default: 11)"},
                        "table_header_bg": {"type": "string", "description": "Table header row background color (hex string)"},
                        "table_header_color": {"type": "string", "description": "Table header text color (hex string)"},
                        "table_alternate_bg": {"type": "string", "description": "Table alternating row background color (hex string)"}
                    }
                }
            },
            "required": ["doc_id", "title", "sections"]
        }
    },
    {
        "name": "edit_word_document",
        "description": "Incrementally edit an existing Word document artifact section-by-section, creating a new version (v2, v3) while preserving history.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "doc_id": {"type": "string", "description": "Document ID slug of the document to edit"},
                "target_heading_or_index": {"type": "string", "description": "Existing heading title or section index to target"},
                "action": {"type": "string", "enum": ["replace", "append", "insert_after", "delete"], "description": "Edit action to execute"},
                "new_content": {
                    "type": "object",
                    "description": "New section object with heading, level, body, or table",
                    "properties": {
                        "heading": {"type": "string"},
                        "level": {"type": "integer"},
                        "body": {"type": "string"},
                        "page_break": {"type": "boolean"},
                        "table": {"type": "object"}
                    }
                }
            },
            "required": ["doc_id", "target_heading_or_index", "action", "new_content"]
        }
    },
    {
        "name": "get_word_document",
        "description": "Retrieve current version state and structure of a Word document artifact.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "doc_id": {"type": "string", "description": "Document ID slug"}
            },
            "required": ["doc_id"]
        }
    },
    {
        "name": "export_word_document",
        "description": "Export the latest version of a Word document artifact as native .docx binary file.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "doc_id": {"type": "string", "description": "Document ID slug"}
            },
            "required": ["doc_id"]
        }
    }
]


async def execute_word_tool(tool_name: str, args: dict) -> str:
    """
    Executes Microsoft Word custom MCP tools and returns JSON response / artifact state.
    """
    clean_tool = tool_name.replace("microsoft_word__", "").replace("word__", "")
    
    if clean_tool == "create_word_document":
        doc_id = args.get("doc_id") or "word_doc_" + str(int(datetime.now().timestamp()))
        title = args.get("title", "Untitled Document")
        sections = args.get("sections", [])
        author = args.get("author")
        header_text = args.get("header_text")
        footer_text = args.get("footer_text", "Page {page} of {pages}")
        style = args.get("style", {})
        
        doc_state = word_engine.create_document(
            doc_id=doc_id,
            title=title,
            sections=sections,
            author=author,
            header_text=header_text,
            footer_text=footer_text,
            style=style
        )
        return json.dumps({
            "status": "success",
            "message": f"Successfully created Word document '{title}' (v1). Artifact filename: '{doc_id}.docx'.",
            "artifact": {
                "name": f"{doc_id}.docx",
                "content": json.dumps(doc_state),
                "type": "word"
            },
            "doc_state": doc_state
        })

    elif clean_tool == "edit_word_document":
        doc_id = args.get("doc_id", "")
        target = args.get("target_heading_or_index", "")
        action = args.get("action", "append")
        new_content = args.get("new_content", {})

        doc_state = word_engine.edit_document(
            doc_id=doc_id,
            target_heading_or_index=target,
            action=action,
            new_content=new_content
        )
        return json.dumps({
            "status": "success",
            "message": f"Successfully updated Word document '{doc_id}' to version v{doc_state['version']}.",
            "artifact": {
                "name": f"{doc_id}.docx",
                "content": json.dumps(doc_state),
                "type": "word"
            },
            "doc_state": doc_state
        })

    elif clean_tool == "get_word_document":
        doc_id = args.get("doc_id", "")
        history = word_engine.documents.get(doc_id, [])
        if not history:
            return json.dumps({"status": "error", "message": f"Word document '{doc_id}' not found."})
        return json.dumps({"status": "success", "versions": history, "latest": history[-1]})

    elif clean_tool == "export_word_document":
        doc_id = args.get("doc_id", "")
        history = word_engine.documents.get(doc_id, [])
        if not history:
            return json.dumps({"status": "error", "message": f"Word document '{doc_id}' not found."})
        latest = history[-1]
        return json.dumps({
            "status": "success",
            "message": f"Ready for export as native .docx file.",
            "doc_id": doc_id,
            "export_url": f"/api/mcp/word/export-docx"
        })

    return json.dumps({"status": "error", "message": f"Unknown Word tool '{clean_tool}'"})
